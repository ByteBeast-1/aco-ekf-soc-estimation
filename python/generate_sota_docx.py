import os
import re
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def markdown_to_docx(md_path, docx_path, title_text):
    if not os.path.exists(md_path):
        print(f"Error: {md_path} does not exist")
        return

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    doc = docx.Document()

    # Set page margins (0.75 in for compact professional engineering layout)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # Document Header Title
    title = doc.add_heading(title_text, level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    sub = doc.add_paragraph("Automotive Model-Based Design (MBD) & Real-World Hardware Verification Manual")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.runs[0].font.italic = True
    sub.runs[0].font.size = Pt(11)
    sub.runs[0].font.color.rgb = RGBColor(90, 100, 110)

    in_code_block = False
    code_lines = []
    in_table = False
    table_lines = []

    def flush_code(lines_list):
        if not lines_list:
            return
        code_p = doc.add_paragraph()
        code_p.paragraph_format.left_indent = Inches(0.2)
        code_p.paragraph_format.space_before = Pt(4)
        code_p.paragraph_format.space_after = Pt(4)
        code_text = "".join(lines_list)
        run = code_p.add_run(code_text)
        run.font.name = 'Consolas'
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(30, 45, 60)

    def flush_table(lines_list):
        if not lines_list:
            return
        rows_data = []
        for l in lines_list:
            if not l.strip() or ('---' in l and '|' in l and not any(c.isalnum() for c in l)):
                continue
            cells = [c.strip() for c in l.strip().strip('|').split('|')]
            if cells:
                rows_data.append(cells)
        if not rows_data:
            return

        cols_count = max(len(r) for r in rows_data)
        table = doc.add_table(rows=len(rows_data), cols=cols_count)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        for row_idx, r_data in enumerate(rows_data):
            row = table.rows[row_idx]
            is_header = (row_idx == 0)
            for c_idx in range(cols_count):
                cell = row.cells[c_idx]
                text = r_data[c_idx] if c_idx < len(r_data) else ""
                cell.text = text
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:
                    r.font.name = 'Calibri'
                    r.font.size = Pt(9)
                    if is_header:
                        r.font.bold = True
                        r.font.color.rgb = RGBColor(255, 255, 255)
                if is_header:
                    set_cell_background(cell, "1F497D")
                else:
                    if row_idx % 2 == 1:
                        set_cell_background(cell, "F2F5F9")
        doc.add_paragraph()

    for raw_line in lines:
        line = raw_line.rstrip('\r\n')

        # Code block handling
        if line.strip().startswith('```'):
            if in_code_block:
                in_code_block = False
                flush_code(code_lines)
                code_lines = []
            else:
                if in_table:
                    in_table = False
                    flush_table(table_lines)
                    table_lines = []
                in_code_block = True
            continue

        if in_code_block:
            code_lines.append(raw_line)
            continue

        # Markdown Table handling
        if line.strip().startswith('|') and line.strip().endswith('|'):
            in_table = True
            table_lines.append(line)
            continue
        else:
            if in_table:
                in_table = False
                flush_table(table_lines)
                table_lines = []

        if not line.strip():
            continue

        if line.startswith('# '):
            h = doc.add_heading(line[2:].strip(), level=1)
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(4)
        elif line.startswith('## '):
            h = doc.add_heading(line[3:].strip(), level=2)
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(3)
        elif line.startswith('### '):
            h = doc.add_heading(line[4:].strip(), level=3)
            h.paragraph_format.space_before = Pt(8)
            h.paragraph_format.space_after = Pt(2)
        elif line.startswith('#### '):
            h = doc.add_heading(line[5:].strip(), level=4)
            h.paragraph_format.space_before = Pt(6)
            h.paragraph_format.space_after = Pt(2)
        elif line.startswith('> '):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            run = p.add_run(line[2:].strip())
            run.font.italic = True
            run.font.color.rgb = RGBColor(70, 80, 95)
        elif line.startswith('* ') or line.startswith('- '):
            p = doc.add_paragraph(style='List Bullet')
            text = line[2:].strip()
            parts = re.split(r'(\*\*.*?\*\*)', text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = p.add_run(part[2:-2])
                    r.bold = True
                else:
                    p.add_run(part)
        else:
            p = doc.add_paragraph()
            parts = re.split(r'(\*\*.*?\*\*)', line)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = p.add_run(part[2:-2])
                    r.bold = True
                else:
                    p.add_run(part)

    if in_code_block:
        flush_code(code_lines)
    if in_table:
        flush_table(table_lines)

    doc.save(docx_path)
    print(f"Successfully generated: {docx_path}")
